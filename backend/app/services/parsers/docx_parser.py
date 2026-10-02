import docx
from typing import Tuple, IO

def extract_text_from_docx(file_stream: IO[bytes]) -> Tuple[str, int]:
    """
    Extracts text from a DOCX file stream using python-docx.
    Returns a tuple of (extracted_text, page_count).
    Since DOCX doesn't have a strict concept of pages like PDF, page_count will be estimated or set to 1.
    """
    try:
        doc = docx.Document(file_stream)
    except Exception as e:
        raise ValueError("Invalid or corrupted DOCX file.")

    extracted_parts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            extracted_parts.append(text)

    # Extract tables (basic text extraction)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    row_text.append(text)
            if row_text:
                extracted_parts.append(" | ".join(row_text))

    full_text = "\n".join(extracted_parts).strip()
    
    # Estimate page count (very roughly, ~3000 chars per page)
    estimated_pages = max(1, len(full_text) // 3000)
    
    return full_text, estimated_pages
